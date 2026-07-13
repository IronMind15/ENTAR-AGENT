"""
Excel 文档提取器

从 Excel 文件中提取结构化记录，用于后续切块和入库。
当前主要处理 PCS 参数表的故障代码数据。
"""

import os
import logging
from typing import Optional

logger = logging.getLogger("doc_mgr.extractors.excel")

# ===== PCS 参数表默认配置 =====
DEFAULT_SHEET = "遥信（DI）"
DEFAULT_START_ROW = 51
DEFAULT_COLUMNS = {
    1: "seq", 2: "name", 3: "address", 4: "bit_address",
    5: "attribute", 6: "data_type", 7: "default_value",
    8: "notes", 9: "description", 10: "notes2",
    11: "fault_code", 12: "cause",
}


def safe_str(v) -> str:
    """安全转字符串，None → '' """
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip()


def extract_excel_rows(
    filepath: str,
    sheet_name: str = DEFAULT_SHEET,
    start_row: int = DEFAULT_START_ROW,
    columns: Optional[dict[int, str]] = None,
) -> list[dict]:
    """从 Excel 中提取结构化记录列表

    Args:
        filepath: Excel 文件路径
        sheet_name: sheet 名称
        start_row: 数据起始行号（1-indexed）
        columns: 列号 → 字段名映射

    Returns:
        记录列表，每条包含 COLUMNS 中定义的所有字段 + _row_num + _sheet_name
    """
    import openpyxl

    cols = columns or DEFAULT_COLUMNS

    if not os.path.exists(filepath):
        logger.error(f"文件不存在: {filepath}")
        return []

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if sheet_name not in wb.sheetnames:
        logger.error(f"Sheet '{sheet_name}' 未找到（已有: {wb.sheetnames}）")
        wb.close()
        return []

    ws = wb[sheet_name]
    records = []

    for row_idx in range(start_row, ws.max_row + 1):
        raw: dict = {}
        has_data = False
        for col, field in cols.items():
            v = ws.cell(row=row_idx, column=col).value
            raw[field] = safe_str(v)
            if raw[field]:
                has_data = True

        if not has_data:
            continue

        raw["_row_num"] = row_idx
        raw["_sheet_name"] = sheet_name
        records.append(raw)

    wb.close()
    logger.info(f"Excel 提取完成: {len(records)} 条记录 ({sheet_name} 第{start_row}行起)")
    return records


def format_excel_row(raw: dict) -> str:
    """将单行 Excel 记录格式化为可搜索文本

    拼接 name / description / cause / notes / fault_code 等关键字段。
    """
    parts = []
    field_labels = {
        "fault_code": "故障代码",
        "name": "名称",
        "description": "说明",
        "cause": "故障原因",
        "notes": "备注",
    }
    for field, label in field_labels.items():
        val = raw.get(field, "")
        if val:
            parts.append(f"{label}：{val}")
    return " | ".join(parts)
