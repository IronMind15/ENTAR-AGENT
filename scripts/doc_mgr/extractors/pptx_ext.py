"""
PPT 文档提取器

使用 python-pptx 提取 .pptx 文件的文本内容，
按 Slide 为单位输出 Markdown 格式，以便复用 MarkdownChunker 进行切块。

转换规则：
  - 每个 Slide → ## Slide N: 标题
  - 正文文本框 → 段落
  - 备注 → > 备注内容（blockquote 格式）
"""

import logging

logger = logging.getLogger("doc_mgr.extractors.pptx")


def extract_pptx_text(filepath: str) -> str:
    """提取 PPT 文档文本，按 Slide 输出 Markdown 格式

    Args:
        filepath: .pptx 文件路径

    Returns:
        Markdown 格式文本；提取失败返回空字符串
    """
    try:
        from pptx import Presentation
    except ImportError:
        logger.error("python-pptx 未安装，请运行 pip install python-pptx")
        return ""

    try:
        prs = Presentation(filepath)
    except Exception as e:
        logger.error(f"无法打开 PPT 文件 {filepath}: {e}")
        return ""

    parts: list[str] = []

    for slide_idx, slide in enumerate(prs.slides, 1):
        slide_title = _get_slide_title(slide)
        header = f"## Slide {slide_idx}: {slide_title}" if slide_title else f"## Slide {slide_idx}"
        parts.append(header)

        # 提取所有文本框内容
        slide_texts = _extract_slide_texts(slide)
        if slide_texts:
            parts.append(slide_texts)

        # 提取备注
        notes = _extract_slide_notes(slide)
        if notes:
            parts.append(notes)

    result = "\n\n".join(parts)
    if not result.strip():
        logger.warning(f"PPT 文件未提取到文本: {filepath}")
    return result


def _get_slide_title(slide) -> str:
    """获取 Slide 标题

    python-pptx 的 slide.shapes.title 会返回标题占位符（如果存在）。
    """
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        return slide.shapes.title.text_frame.text.strip()
    return ""


def _extract_slide_texts(slide) -> str:
    """提取 Slide 中所有文本框的文本（跳过标题占位符，避免重复）"""
    texts: list[str] = []
    title_shape_id = slide.shapes.title.shape_id if slide.shapes.title else None

    for shape in slide.shapes:
        # 跳过标题占位符（已在 header 中输出）
        if shape.shape_id == title_shape_id:
            continue

        if not shape.has_text_frame:
            continue

        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                texts.append(text)

    return "\n\n".join(texts)


def _extract_slide_notes(slide) -> str:
    """提取 Slide 的备注内容"""
    if not slide.has_notes_slide:
        return ""

    notes_slide = slide.notes_slide
    if not notes_slide.notes_text_frame:
        return ""

    notes_text = notes_slide.notes_text_frame.text.strip()
    if not notes_text:
        return ""

    # 用 blockquote 格式输出备注
    lines = notes_text.split("\n")
    return "\n".join(f"> {line}" for line in lines)
