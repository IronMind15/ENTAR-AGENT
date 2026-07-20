"""将恩特小助手设计说明.docx 转换为 Markdown 格式"""
import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/恩特小助手设计说明.docx')

def get_heading_level(style_name):
    if 'Heading 1' in style_name:
        return 1
    elif 'Heading 2' in style_name:
        return 2
    elif 'Heading 3' in style_name:
        return 3
    return None


def extract_table(table, table_index=None):
    """提取表格为 Markdown"""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        rows.append(cells)
    if not rows:
        return ''

    max_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_cols:
            r.append('')

    # 单行单列 = 架构图
    if max_cols == 1 and len(rows) == 1:
        return rows[0][0]

    md_lines = []
    md_lines.append('| ' + ' | '.join(rows[0]) + ' |')
    md_lines.append('| ' + ' | '.join(['---'] * max_cols) + ' |')
    for row in rows[1:]:
        if all(c == '' for c in row):
            continue
        md_lines.append('| ' + ' | '.join(row) + ' |')
    return '\n'.join(md_lines)


# ===== 遍历 body 子元素，按顺序合并段落和表格 =====
body = doc.element.body

all_paragraphs = doc.paragraphs   # list of Paragraph objects
all_tables = doc.tables           # list of Table objects

p_idx = 0
t_idx = 0

output_lines = []
in_toc = False

for child in body:
    tag = child.tag.split('}')[-1]

    if tag == 'p':
        if p_idx >= len(all_paragraphs):
            continue
        para = all_paragraphs[p_idx]
        p_idx += 1

        text = para.text
        style = para.style.name

        if not text.strip():
            continue

        # --- 修订记录 ---
        if text.strip() == '修订记录':
            output_lines.append('\n---\n')
            output_lines.append('## 修订记录\n')
            continue

        # --- 目录 ---
        if text.strip() == '目录':
            in_toc = True
            continue

        # --- 目录中的条目跳过 ---
        if in_toc:
            # 遇到第一个正文 Heading 1 时退出目录模式
            if style == 'Heading 1':
                in_toc = False
                # 继续往下，处理这个标题
            else:
                continue

        # --- 标题 ---
        level = get_heading_level(style)
        if level:
            if level == 1:
                output_lines.append('\n---\n')
            output_lines.append(f"{'#' * level} {text.strip()}\n")
            continue

        # --- 列表 ---
        if 'List' in style or 'Bullet' in style:
            output_lines.append(f'- {text.strip()}\n')
            continue

        # --- 代码 (No Spacing) ---
        if style == 'No Spacing':
            output_lines.append(f'```\n{text}\n```\n')
            continue

        # --- 普通段落 ---
        if style == 'Normal':
            output_lines.append(text.strip() + '\n\n')
            continue

        output_lines.append(text.strip() + '\n\n')

    elif tag == 'tbl':
        if t_idx >= len(all_tables):
            continue
        table = all_tables[t_idx]
        t_idx += 1
        md_table = extract_table(table)
        output_lines.append(md_table + '\n\n')

result = ''.join(output_lines)
result = re.sub(r'\n{4,}', '\n\n\n', result)
result = result.strip()

output_path = 'docs/恩特小助手设计说明.md'
with open(output_path, 'w', encoding='utf-8') as f:
    f.write(result)

print(f'✅ 转换完成！输出文件：{output_path}')
print(f'📊 总字符数：{len(result)}')
