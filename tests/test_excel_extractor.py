"""Excel 提取器补测（v1.13.4 审查补：word/pptx/csv 有测试、excel 提取器零覆盖）"""

import os
import shutil
import tempfile
import unittest

from scripts.doc_mgr.extractors.excel import (
    extract_excel_rows, format_excel_row, safe_str,
    DEFAULT_SHEET, DEFAULT_START_ROW, DEFAULT_COLUMNS,
)


class SafeStrTests(unittest.TestCase):
    """safe_str：None/数值/字符串安全转文本"""

    def test_none_to_empty(self):
        self.assertEqual(safe_str(None), "")

    def test_int_and_float(self):
        self.assertEqual(safe_str(42), "42")
        self.assertEqual(safe_str(3.0), "3")
        self.assertEqual(safe_str(3.14), "3.14")

    def test_str_stripped(self):
        self.assertEqual(safe_str("  abc  "), "abc")


class ExtractExcelRowsTests(unittest.TestCase):
    """extract_excel_rows：真实临时 xlsx 提取 + 空行跳过 + 异常路径"""

    def setUp(self):
        import openpyxl
        self._tmp = tempfile.mkdtemp(prefix="excel_extractor_")
        self._path = os.path.join(self._tmp, "t.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = DEFAULT_SHEET
        # 在 DEFAULT_START_ROW 行写一行故障数据（模拟 PCS 参数表结构）
        vals = {1: "1", 2: "急停告警", 3: "0x3004", 4: "bit0", 5: "R",
                6: "bit", 7: "0", 8: "一级", 9: "1-告警", 10: "1",
                11: "d4-1", 12: "外部急停信号闭合"}
        for col, v in vals.items():
            ws.cell(row=DEFAULT_START_ROW, column=col, value=v)
        wb.save(self._path)
        wb.close()
        self.addCleanup(shutil.rmtree, self._tmp, True)

    def test_extract_rows_parses_fields(self):
        records = extract_excel_rows(self._path)
        self.assertEqual(len(records), 1)
        r = records[0]
        self.assertEqual(r["fault_code"], "d4-1")
        self.assertEqual(r["name"], "急停告警")
        self.assertEqual(r["cause"], "外部急停信号闭合")
        self.assertEqual(r["_row_num"], DEFAULT_START_ROW)
        self.assertEqual(r["_sheet_name"], DEFAULT_SHEET)

    def test_extract_skips_empty_rows(self):
        """全空行被跳过（不产生记录）"""
        import openpyxl
        wb = openpyxl.load_workbook(self._path)
        ws = wb[DEFAULT_SHEET]
        for col in DEFAULT_COLUMNS:
            ws.cell(row=DEFAULT_START_ROW + 1, column=col, value=None)
        wb.save(self._path)
        wb.close()
        records = extract_excel_rows(self._path)
        self.assertEqual(len(records), 1)

    def test_missing_sheet_returns_empty(self):
        records = extract_excel_rows(self._path, sheet_name="不存在的sheet")
        self.assertEqual(records, [])

    def test_missing_file_returns_empty(self):
        records = extract_excel_rows(os.path.join(self._tmp, "nope.xlsx"))
        self.assertEqual(records, [])


class FormatExcelRowTests(unittest.TestCase):
    """format_excel_row：关键字段带中文标签拼接"""

    def test_formats_key_fields(self):
        raw = {"fault_code": "d4-1", "name": "急停告警",
               "description": "1-告警", "cause": "外部急停信号闭合"}
        s = format_excel_row(raw)
        self.assertIn("故障代码：d4-1", s)
        self.assertIn("名称：急停告警", s)
        self.assertIn("故障原因：外部急停信号闭合", s)

    def test_empty_raw(self):
        self.assertEqual(format_excel_row({}), "")


if __name__ == "__main__":
    unittest.main()
