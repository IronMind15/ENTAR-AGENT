"""
新文档格式提取器测试

覆盖 Word (.docx)、PPT (.pptx)、CSV (.csv) 三种格式的：
  1. 提取器单元测试（直接调用 extract 函数）
  2. 引擎集成测试（process_file 路由 + 处理函数）
"""

import importlib
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = PROJECT_ROOT / "scripts"

from scripts.doc_mgr import engine
from scripts.doc_mgr.models import Document


class FakeStore:
    def __init__(self, existing=None):
        self.add_calls = []
        self.existing = existing or {"ids": [], "metadatas": []}

    def get(self, collection, ids=None, where=None):
        return self.existing

    def add(self, collection, ids, documents, metadatas):
        self.add_calls.append({
            "collection": collection,
            "ids": list(ids),
            "documents": list(documents),
            "metadatas": list(metadatas),
        })
        return len(ids)

    def replace_document(self, collection, doc_id, file_name,
                         ids, documents, metadatas, legacy_ids=None):
        self.add_calls.append({
            "collection": collection,
            "doc_id": doc_id,
            "file_name": file_name,
            "ids": list(ids),
            "documents": list(documents),
            "metadatas": list(metadatas),
            "legacy_ids": list(legacy_ids or []),
        })
        return len(ids)


# ==================== Word 提取器测试 ====================

@unittest.skipUnless(
    importlib.util.find_spec("docx"), "python-docx 未安装"
)
class WordExtractorTests(unittest.TestCase):

    def test_extract_headings_and_paragraphs(self):
        """测试标题和普通段落提取"""
        from docx import Document as DocxDocument
        from scripts.doc_mgr.extractors.word import extract_docx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "test.docx"
            doc = DocxDocument()
            doc.add_heading("第一章 概述", level=1)
            doc.add_paragraph("这是正文内容。")
            doc.add_heading("1.1 背景", level=2)
            doc.add_paragraph("背景描述。")
            doc.save(str(file_path))

            text = extract_docx_text(str(file_path))

            self.assertIn("# 第一章 概述", text)
            self.assertIn("这是正文内容。", text)
            self.assertIn("## 1.1 背景", text)
            self.assertIn("背景描述。", text)

    def test_extract_table_as_markdown(self):
        """测试表格转为 Markdown 格式"""
        from docx import Document as DocxDocument
        from scripts.doc_mgr.extractors.word import extract_docx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "table.docx"
            doc = DocxDocument()
            doc.add_heading("数据表", level=1)
            table = doc.add_table(rows=3, cols=3)
            # 表头
            table.rows[0].cells[0].text = "名称"
            table.rows[0].cells[1].text = "数值"
            table.rows[0].cells[2].text = "单位"
            # 数据行
            table.rows[1].cells[0].text = "电压"
            table.rows[1].cells[1].text = "220"
            table.rows[1].cells[2].text = "V"
            table.rows[2].cells[0].text = "电流"
            table.rows[2].cells[1].text = "10"
            table.rows[2].cells[2].text = "A"
            doc.save(str(file_path))

            text = extract_docx_text(str(file_path))

            self.assertIn("| 名称 | 数值 | 单位 |", text)
            self.assertIn("| --- | --- | --- |", text)
            self.assertIn("| 电压 | 220 | V |", text)
            self.assertIn("| 电流 | 10 | A |", text)

    def test_empty_docx_returns_empty(self):
        """空文档返回空字符串"""
        from docx import Document as DocxDocument
        from scripts.doc_mgr.extractors.word import extract_docx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "empty.docx"
            doc = DocxDocument()
            doc.save(str(file_path))

            text = extract_docx_text(str(file_path))
            self.assertEqual(text, "")

    def test_non_docx_file_returns_empty(self):
        """非 docx 文件返回空字符串"""
        from scripts.doc_mgr.extractors.word import extract_docx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "test.txt"
            file_path.write_text("不是 docx 文件", encoding="utf-8")

            text = extract_docx_text(str(file_path))
            self.assertEqual(text, "")


# ==================== PPT 提取器测试 ====================

@unittest.skipUnless(
    importlib.util.find_spec("pptx"), "python-pptx 未安装"
)
class PptxExtractorTests(unittest.TestCase):

    def test_extract_slides_with_notes(self):
        """测试按 Slide 提取文本和备注"""
        from pptx import Presentation as PptxPresentation
        from scripts.doc_mgr.extractors.pptx_ext import extract_pptx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "test.pptx"
            prs = PptxPresentation()

            # Slide 1：带标题和正文
            slide1 = prs.slides.add_slide(prs.slide_layouts[1])
            slide1.shapes.title.text = "项目概述"
            slide1.placeholders[1].text = "这是项目简介内容。"

            # Slide 2：带备注
            slide2 = prs.slides.add_slide(prs.slide_layouts[1])
            slide2.shapes.title.text = "技术架构"
            slide2.placeholders[1].text = "系统采用 RAG 架构。"
            notes = slide2.notes_slide
            notes.notes_text_frame.text = "备注：重点讲解检索部分"

            prs.save(str(file_path))

            text = extract_pptx_text(str(file_path))

            self.assertIn("## Slide 1: 项目概述", text)
            self.assertIn("这是项目简介内容。", text)
            self.assertIn("## Slide 2: 技术架构", text)
            self.assertIn("系统采用 RAG 架构。", text)
            self.assertIn("> 备注：重点讲解检索部分", text)

    def test_empty_pptx_returns_empty(self):
        """空 PPT 返回空字符串"""
        from pptx import Presentation as PptxPresentation
        from scripts.doc_mgr.extractors.pptx_ext import extract_pptx_text

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "empty.pptx"
            prs = PptxPresentation()
            prs.save(str(file_path))

            text = extract_pptx_text(str(file_path))
            self.assertEqual(text, "")


# ==================== CSV 提取器测试 ====================

class CsvExtractorTests(unittest.TestCase):

    def test_extract_utf8_csv(self):
        """测试 UTF-8 CSV 提取"""
        from scripts.doc_mgr.extractors.csv_ext import extract_csv_rows, format_csv_row

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "test.csv"
            file_path.write_text(
                "名称,数值,单位\n电压,220,V\n电流,10,A\n",
                encoding="utf-8",
            )

            records = extract_csv_rows(str(file_path))

            self.assertEqual(len(records), 2)
            self.assertEqual(records[0]["名称"], "电压")
            self.assertEqual(records[0]["数值"], "220")
            self.assertEqual(records[0]["_row_num"], 2)
            self.assertEqual(records[1]["名称"], "电流")

            # 测试格式化
            text = format_csv_row(records[0])
            self.assertIn("名称：电压", text)
            self.assertIn("数值：220", text)

    def test_extract_gbk_csv(self):
        """测试 GBK 编码 CSV（中文 Windows 常见）"""
        from scripts.doc_mgr.extractors.csv_ext import extract_csv_rows

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "gbk.csv"
            file_path.write_text(
                "名称,数值\n电压,220\n",
                encoding="gbk",
            )

            records = extract_csv_rows(str(file_path))

            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]["名称"], "电压")

    def test_empty_csv_returns_empty(self):
        """空 CSV 返回空列表"""
        from scripts.doc_mgr.extractors.csv_ext import extract_csv_rows

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "empty.csv"
            file_path.write_text("", encoding="utf-8")

            records = extract_csv_rows(str(file_path))
            self.assertEqual(records, [])

    def test_csv_with_empty_rows_skipped(self):
        """空行自动跳过"""
        from scripts.doc_mgr.extractors.csv_ext import extract_csv_rows

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "with_blanks.csv"
            file_path.write_text(
                "名称,数值\n电压,220\n,\n电流,10\n",
                encoding="utf-8",
            )

            records = extract_csv_rows(str(file_path))
            # 中间的空行应该被跳过
            self.assertEqual(len(records), 2)


# ==================== 引擎集成测试 ====================

class EngineRoutingTests(unittest.TestCase):
    """测试 process_file 对新格式的路由"""

    def test_docx_routes_to_process_word(self):
        """docx 文件路由到 _process_word"""
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "sample.docx"
            file_path.write_bytes(b"fake-docx")
            processor = Mock(return_value=Document(
                file_name="sample.docx",
                file_path=str(file_path),
                collection="standards",
                status="done",
            ))
            with patch.object(engine, "get_store", return_value=FakeStore()), \
                    patch.object(engine, "_process_word", processor), \
                    patch.object(engine, "_record_sync_status"):
                doc = engine.process_file(str(file_path))

            self.assertEqual("standards", doc.collection)
            processor.assert_called_once()

    def test_pptx_routes_to_process_pptx(self):
        """pptx 文件路由到 _process_pptx"""
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "slides.pptx"
            file_path.write_bytes(b"fake-pptx")
            processor = Mock(return_value=Document(
                file_name="slides.pptx",
                file_path=str(file_path),
                collection="standards",
                status="done",
            ))
            with patch.object(engine, "get_store", return_value=FakeStore()), \
                    patch.object(engine, "_process_pptx", processor), \
                    patch.object(engine, "_record_sync_status"):
                doc = engine.process_file(str(file_path))

            self.assertEqual("standards", doc.collection)
            processor.assert_called_once()

    def test_csv_routes_to_process_csv(self):
        """csv 文件路由到 _process_csv"""
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "data.csv"
            file_path.write_text("name,value\na,1\n", encoding="utf-8")
            processor = Mock(return_value=Document(
                file_name="data.csv",
                file_path=str(file_path),
                collection="error_codes",
                status="done",
            ))
            with patch.object(engine, "get_store", return_value=FakeStore()), \
                    patch.object(engine, "_process_csv", processor), \
                    patch.object(engine, "_record_sync_status"):
                doc = engine.process_file(str(file_path))

            self.assertEqual("error_codes", doc.collection)
            processor.assert_called_once()

    def test_unsupported_ext_returns_error(self):
        """不支持的格式返回 error"""
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "test.xyz"
            file_path.write_bytes(b"unknown")

            with patch.object(engine, "_record_sync_status"):
                doc = engine.process_file(str(file_path))

            self.assertEqual("error", doc.status)
            self.assertIn("不支持", doc.message)


if __name__ == "__main__":
    unittest.main()
